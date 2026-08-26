"""Generate a sample DOCX report for local experimentation."""

import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_shading(cell, color_hex):
    """Apply a background fill to one Word table cell."""
    shading_xml = f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>'
    cell._tc.get_or_add_tcPr().append(parse_xml(shading_xml))

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Set the internal margins of one Word table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def create_report():
    """Create the standalone sample DOCX report in the sandbox workspace."""
    # This standalone example demonstrates Word formatting and is separate
    # from the Markdown/HTML report path used by the main application.
    doc = Document()
    
    # Page setup
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Styles
    styles = doc.styles
    normal_style = styles['Normal']
    normal_style.font.name = 'Arial'
    normal_style.font.size = Pt(10.5)
    normal_style.font.color.rgb = RGBColor(51, 51, 51)
    normal_style.paragraph_format.line_spacing = 1.15
    normal_style.paragraph_format.space_after = Pt(6)

    # Cover Page / Header
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(24)
    title_p.paragraph_format.space_after = Pt(4)
    title_run = title_p.add_run("STUDENT PERFORMANCE ANALYTICAL REPORT")
    title_run.font.size = Pt(22)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(31, 78, 121) # Navy

    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_after = Pt(24)
    sub_run = sub_p.add_run("A Comprehensive Examination of Determinants Driving Student Academic Success")
    sub_run.font.size = Pt(13)
    sub_run.font.italic = True
    sub_run.font.color.rgb = RGBColor(89, 89, 89)

    # Divider line
    rule = doc.add_paragraph()
    rule.paragraph_format.space_after = Pt(18)
    r_run = rule.add_run("―" * 55)
    r_run.font.color.rgb = RGBColor(31, 78, 121)
    r_run.font.bold = True

    # Section 1: Executive Summary
    h1 = doc.add_paragraph()
    h1.paragraph_format.space_before = Pt(16)
    h1.paragraph_format.space_after = Pt(6)
    h1_run = h1.add_run("1. Executive Summary")
    h1_run.font.size = Pt(14)
    h1_run.font.bold = True
    h1_run.font.color.rgb = RGBColor(31, 78, 121)

    p1 = doc.add_paragraph(
        "This report provides an in-depth empirical analysis of student academic performance based on a rigorous dataset of "
        "10,000 student records. The primary objective is to identify, quantify, and evaluate the key drivers influencing the "
        "Performance Index—ranging from historical academic preparation and study habits to lifestyle factors such as sleep and "
        "extracurricular engagement. Statistical modeling reveals that historical academic achievement ('Previous Scores') is by far "
        "the single most dominant predictor of future performance, exhibiting an exceptionally strong positive correlation (r = 0.915). "
        "Additionally, dedicated study hours serve as a powerful secondary catalyst for academic excellence. Conversely, lifestyle variables "
        "such as sleep duration and extracurricular involvement exhibit negligible direct correlation with overall academic outcomes."
    )

    # Section 2: Dataset Architecture & Quality
    h2 = doc.add_paragraph()
    h2.paragraph_format.space_before = Pt(16)
    h2.paragraph_format.space_after = Pt(6)
    h2_run = h2.add_run("2. Dataset Architecture & Data Quality Audit")
    h2_run.font.size = Pt(14)
    h2_run.font.bold = True
    h2_run.font.color.rgb = RGBColor(31, 78, 121)

    doc.add_paragraph(
        "The evaluated dataset comprises 10,000 rows and 6 distinct structural columns. A comprehensive data quality audit was "
        "conducted prior to statistical aggregation:"
    )

    bullet1 = doc.add_paragraph(style='List Bullet')
    r1 = bullet1.add_run("Total Records: ")
    r1.font.bold = True
    bullet1.add_run("10,000 student observations across 6 quantitative and categorical attributes.")

    bullet2 = doc.add_paragraph(style='List Bullet')
    r2 = bullet2.add_run("Missing Values: ")
    r2.font.bold = True
    bullet2.add_run("Zero missing values were detected across all 10,000 rows, ensuring 100% data completeness.")

    bullet3 = doc.add_paragraph(style='List Bullet')
    r3 = bullet3.add_run("Duplicate Records: ")
    r3.font.bold = True
    bullet3.add_run("127 duplicate rows were identified. These represent identical student profiles and were retained as they do not skew underlying distributional parameters or correlation matrices.")

    # Table of Dataset Columns
    table = doc.add_table(rows=7, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Column Name", "Data Type", "Mean / Distribution", "Description"]
    col_widths = [Inches(1.8), Inches(1.1), Inches(1.8), Inches(2.3)]

    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        set_cell_shading(hdr_cells[i], "1F4E79")
        set_cell_margins(hdr_cells[i], 120, 120, 150, 150)
        for p in hdr_cells[i].paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for run in p.runs:
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
                run.font.size = Pt(9.5)

    data_rows = [
        ("Hours Studied", "Integer", "Mean: 4.99 (Min: 1, Max: 9)", "Total daily study hours recorded"),
        ("Previous Scores", "Integer", "Mean: 69.45 (Min: 40, Max: 99)", "Scores achieved in previous academic exams"),
        ("Extracurricular Activities", "Category", "No: 5,052 | Yes: 4,948", "Participation in extracurriculars"),
        ("Sleep Hours", "Integer", "Mean: 6.53 (Min: 4, Max: 9)", "Average daily hours of sleep"),
        ("Sample Question Papers", "Integer", "Mean: 4.58 (Min: 0, Max: 9)", "Number of practice question papers solved"),
        ("Performance Index", "Float", "Mean: 55.22 (Min: 10, Max: 100)", "Overall academic performance score")
    ]

    for row_idx, row_data in enumerate(data_rows, start=1):
        row_cells = table.rows[row_idx].cells
        bg_color = "F2F5F8" if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, cell_value in enumerate(row_data):
            row_cells[col_idx].text = cell_value
            set_cell_shading(row_cells[col_idx], bg_color)
            set_cell_margins(row_cells[col_idx], 100, 100, 150, 150)
            for p in row_cells[col_idx].paragraphs:
                for run in p.runs:
                    run.font.size = Pt(9)
                    run.font.color.rgb = RGBColor(51, 51, 51)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # Section 3: Statistical Summary & Distribution Analysis
    h3 = doc.add_paragraph()
    h3.paragraph_format.space_before = Pt(16)
    h3.paragraph_format.space_after = Pt(6)
    h3_run = h3.add_run("3. Statistical Summary & Distribution Analysis")
    h3_run.font.size = Pt(14)
    h3_run.font.bold = True
    h3_run.font.color.rgb = RGBColor(31, 78, 121)

    doc.add_paragraph(
        "A rigorous statistical evaluation of the dataset highlights balanced numerical distributions across key operational metrics. "
        "The Performance Index exhibits a symmetric spread spanning from a minimum of 10.0 to a maximum of 100.0, with a mean "
        "of 55.22 and a median of 55.0. Standard deviations across hours studied (2.59), previous scores (17.34), and sleep hours (1.70) "
        "indicate healthy sample variance without artificial clustering."
    )

    # Insert Chart
    doc.add_paragraph().paragraph_format.space_after = Pt(6)
    try:
        doc.add_picture("workspace/charts/chart_1a4283cf.png", width=Inches(6.0))
        img_p = doc.paragraphs[-1]
        img_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap_p = doc.add_paragraph()
        cap_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap_run = cap_p.add_run("Figure 1: Distribution of Student Performance Index (10,000 observations)")
        cap_run.font.size = Pt(9)
        cap_run.font.italic = True
        cap_run.font.color.rgb = RGBColor(89, 89, 89)
        cap_p.paragraph_format.space_after = Pt(16)
    except Exception as e:
        pass

    # Section 4: Correlation Matrix & Key Performance Drivers
    h4 = doc.add_paragraph()
    h4.paragraph_format.space_before = Pt(16)
    h4.paragraph_format.space_after = Pt(6)
    h4_run = h4.add_run("4. Correlation Matrix & Key Performance Drivers")
    h4_run.font.size = Pt(14)
    h4_run.font.bold = True
    h4_run.font.color.rgb = RGBColor(31, 78, 121)

    doc.add_paragraph(
        "To uncover the primary determinants of academic success, a full Pearson correlation matrix was computed across all numeric "
        "variables. The findings confirm that academic outcome is heavily dictated by foundational preparation and active study commitment."
    )

    # Correlation Table
    corr_table = doc.add_table(rows=6, cols=6)
    corr_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    corr_headers = ["Metric", "Hours Studied", "Previous Scores", "Sleep Hours", "Sample Papers", "Performance Index"]
    
    corr_hdr_cells = corr_table.rows[0].cells
    for i, title in enumerate(corr_headers):
        corr_hdr_cells[i].text = title
        set_cell_shading(corr_hdr_cells[i], "1F4E79")
        set_cell_margins(corr_hdr_cells[i], 100, 100, 100, 100)
        for p in corr_hdr_cells[i].paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
                run.font.size = Pt(8.5)

    corr_data = [
        ["Hours Studied", "1.000", "-0.012", "0.001", "0.018", "0.374"],
        ["Previous Scores", "-0.012", "1.000", "0.006", "0.008", "0.915"],
        ["Sleep Hours", "0.001", "0.006", "1.000", "0.004", "0.048"],
        ["Sample Papers", "0.018", "0.008", "0.004", "1.000", "0.043"],
        ["Performance Index", "0.374", "0.915", "0.048", "0.043", "1.000"]
    ]

    for row_idx, row_data in enumerate(corr_data, start=1):
        row_cells = corr_table.rows[row_idx].cells
        bg_color = "F2F5F8" if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, cell_value in enumerate(row_data):
            row_cells[col_idx].text = cell_value
            set_cell_shading(row_cells[col_idx], bg_color)
            set_cell_margins(row_cells[col_idx], 80, 80, 100, 100)
            for p in row_cells[col_idx].paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if col_idx > 0 else WD_ALIGN_PARAGRAPH.LEFT
                for run in p.runs:
                    run.font.size = Pt(8.5)
                    run.font.color.rgb = RGBColor(51, 51, 51)
                    if col_idx == 5 or (col_idx == row_idx and col_idx < 5):
                        run.font.bold = True

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    doc.add_paragraph(
        "Key takeaways from the correlation analysis include:"
    )

    b_corr1 = doc.add_paragraph(style='List Bullet')
    rb1 = b_corr1.add_run("Previous Scores (r = 0.915): ")
    rb1.font.bold = True
    b_corr1.add_run("Demonstrates an extraordinarily high positive correlation with the Performance Index. Academic competency is highly persistent; students who perform well historically consistently maintain superior performance.")

    b_corr2 = doc.add_paragraph(style='List Bullet')
    rb2 = b_corr2.add_run("Hours Studied (r = 0.374): ")
    rb2.font.bold = True
    b_corr2.add_run("Represents a robust secondary driver of academic achievement. Increased dedicated study time directly translates into higher performance indices.")

    b_corr3 = doc.add_paragraph(style='List Bullet')
    rb3 = b_corr3.add_run("Lifestyle & Practice Factors (r < 0.05): ")
    rb3.font.bold = True
    b_corr3.add_run("Sleep hours (r = 0.048) and sample question papers practiced (r = 0.043) show very weak linear associations with overall performance within the observed ranges, indicating that baseline intelligence and consistent study habits dwarf marginal variations in practice volume.")

    # Section 5: Extracurricular Impact Analysis
    h5 = doc.add_paragraph()
    h5.paragraph_format.space_before = Pt(16)
    h5.paragraph_format.space_after = Pt(6)
    h5_run = h5.add_run("5. Extracurricular Engagement Impact Analysis")
    h5_run.font.size = Pt(14)
    h5_run.font.bold = True
    h5_run.font.color.rgb = RGBColor(31, 78, 121)

    doc.add_paragraph(
        "An evaluation of categorical participation in extracurricular activities indicates an even split across the student body: "
        "5,052 students (50.5%) do not participate, while 4,948 students (49.5%) participate. Comparing academic performance across "
        "these cohorts reveals nearly identical performance outcomes:"
    )

    extra_table = doc.add_table(rows=3, cols=5)
    extra_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    extra_headers = ["Extracurricular Status", "Student Count", "Mean Performance", "Median Performance", "Max Performance"]
    
    ex_hdr_cells = extra_table.rows[0].cells
    for i, title in enumerate(extra_headers):
        ex_hdr_cells[i].text = title
        set_cell_shading(ex_hdr_cells[i], "1F4E79")
        set_cell_margins(ex_hdr_cells[i], 100, 100, 100, 100)
        for p in ex_hdr_cells[i].paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
                run.font.size = Pt(9)

    ex_data = [
        ["No (Non-Participants)", "5,052", "54.76", "55.0", "99.0"],
        ["Yes (Participants)", "4,948", "55.70", "55.0", "100.0"]
    ]

    for row_idx, row_data in enumerate(ex_data, start=1):
        row_cells = extra_table.rows[row_idx].cells
        bg_color = "F2F5F8" if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, cell_value in enumerate(row_data):
            row_cells[col_idx].text = cell_value
            set_cell_shading(row_cells[col_idx], bg_color)
            set_cell_margins(row_cells[col_idx], 80, 80, 100, 100)
            for p in row_cells[col_idx].paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if col_idx > 0 else WD_ALIGN_PARAGRAPH.LEFT
                for run in p.runs:
                    run.font.size = Pt(9)
                    run.font.color.rgb = RGBColor(51, 51, 51)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    doc.add_paragraph(
        "Insight: Participation in extracurricular activities does not penalize academic performance. In fact, students engaged "
        "in extracurriculars exhibit a marginally higher mean performance index (55.70 vs. 54.76), demonstrating that well-rounded "
        "lifestyle engagement is fully compatible with strong academic standing."
    )

    # Section 6: Recommendations & Conclusion
    h6 = doc.add_paragraph()
    h6.paragraph_format.space_before = Pt(16)
    h6.paragraph_format.space_after = Pt(6)
    h6_run = h6.add_run("6. Strategic Recommendations & Conclusion")
    h6_run.font.size = Pt(14)
    h6_run.font.bold = True
    h6_run.font.color.rgb = RGBColor(31, 78, 121)

    doc.add_paragraph(
        "Based on empirical findings from the 10,000 student dataset, educational institutions and student counselors should implement "
        "the following targeted strategies:"
    )

    rec1 = doc.add_paragraph(style='List Bullet')
    rr1 = rec1.add_run("Early Intervention Based on Previous Scores: ")
    rr1.font.bold = True
    rec1.add_run("Because historical scores serve as the primary predictor of future success (r = 0.915), struggling students should be identified immediately upon enrollment and provided with remedial academic support.")

    rec2 = doc.add_paragraph(style='List Bullet')
    rr2 = rec2.add_run("Promote Structured Study Habits: ")
    rr2.font.bold = True
    rec2.add_run("With study hours demonstrating a strong positive effect (r = 0.374), academic programs should guide students in establishing disciplined daily study routines rather than relying solely on last-minute preparation.")

    rec3 = doc.add_paragraph(style='List Bullet')
    rr3 = rec3.add_run("Encourage Extracurricular Involvement: ")
    rr3.font.bold = True
    rec3.add_run("Evidence disproves the notion that extracurricular activities detract from academic performance. Institutions should actively encourage student participation in non-academic clubs and sports to promote holistic development.")

    # Save document
    doc.save("reports/Student_Performance_Analysis_Report.docx")
    print("Report successfully saved to reports/Student_Performance_Analysis_Report.docx")

if __name__ == "__main__":
    create_report()
